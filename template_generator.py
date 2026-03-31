from __future__ import annotations

from pathlib import Path
import logging
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

import config


DEFAULT_TEMPLATE_FILENAME = "Class_Define_Template.xlsx"

ITEM_CODE_DB_HEADERS = ["Item_Code", "Item_Name", "Group"]
ITEM_CODE_DB_DEFAULT_ROWS = [
    ("P", "PIPE", "Pipe_Group"),
    ("JN", "NIPPLE", "Pipe_Group"),
    ("JNP", "NIPPLE", "Pipe_Group"),
    ("JN1", "NIPPLE", "Pipe_Group"),
    ("JNP1", "NIPPLE", "Pipe_Group"),
    ("RC", "REDUCER CON", "Fitting_Group"),
    ("RE", "REDUCER ECC", "Fitting_Group"),
    ("RCS", "SWAGE CON", "Fitting_Group"),
    ("RES", "SWAGE ECC", "Fitting_Group"),
    ("E", "ELBOW 90 DEG LR", "Fitting_Group"),
    ("ES", "ELBOW 90 DEG SR", "Fitting_Group"),
    ("E4", "ELBOW 45 DEG LR", "Fitting_Group"),
    ("ES4", "ELBOW 45 DEG SR", "Fitting_Group"),
]

CLASS_DEFINE_HEADERS = [
    "Revision_No",
    "Class_Name",
    "Design_Code",
    "Class_Base_Material",
    "Class_Rating",
    "Corrosion_Allowance",
    "Design_Temperature_From",
    "Design_Temperature_To",
    "Design_Pressure_From",
    "Design_Pressure_To",
    "Fluid_Service",
    "Branch_Table_1",
    "Branch_Table_2",
    "Reducing_Table_1",
    "Reducing_Table_2",
    "Global_Special_Req",
    "Remarks",
]

FLUID_SERVICE_HEADERS = [
    "Class_Name",
    "Fluid_Service_Code",
    "Fluid_Service_Name",
    "Min_Design_Temperature",
    "Max_Design_Temperature",
    "Min_Design_Pressure",
    "Max_Design_Pressure",
    "NDE",
    "PWHT",
]

JOINT_HEADERS = [
    "Class_Name",
    "Size_From",
    "Size_To",
    "Pipe_Joint_Type",
    "Maintenance_Joint_Type",
    "Remarks",
]

SCHEDULE_HEADERS = [
    "Class_Name",
    "Size_From",
    "Size_To",
    "Schedule",
]

REDUCING_TABLE_HEADERS = ["Table_Code", "Size1", "Size2", "Item_Type", "Remarks"]

PIPE_HEADERS = [
    "Class_Name",
    "Item_Code",
    "Size_From",
    "Size_To",
    "Mat_Code",
    "Mat_Class",
    "Manufacturing_Method",
    "End_Type",
    "Dim_Standard",
    "Remarks",
]

FITTING_HEADERS = [
    "Class_Name",
    "Item_Code",
    "Size1_From",
    "Size1_To",
    "Size2_From",
    "Size2_To",
    "Mat_Code",
    "Mat_Class",
    "Manufacturing_Method",
    "Rating",
    "End_Type_1",
    "End_Type_2",
    "Dim_Standard",
    "Remarks",
]

FLANGE_HEADERS = [
    "Class_Name",
    "Item_Code",
    "Size1_From",
    "Size1_To",
    "Size2_From",
    "Size2_To",
    "Mat_Code",
    "Mat_Class",
    "Rating",
    "Facing",
    "End_Type",
    "Bore_Schedule",
    "Dim_Standard",
    "Remarks",
]

VALVE_HEADERS = [
    "Class_Name",
    "Item_Code",
    "Valve_Type",
    "Size_From",
    "Size_To",
    "Body_Mat",
    "Trim_Mat",
    "Rating",
    "End_Type",
    "Operation",
    "Bonnet_Type",
    "Valve_Feature",
    "Dim_Standard",
    "Remarks",
]

HEADER_FONT = Font(bold=True)
HEADER_ALIGNMENT = Alignment(horizontal="center", vertical="center", wrap_text=True)
FREEZE_PANES = "A2"


def _get_logger() -> logging.Logger:
    # Keep module self-contained; if project has a logger module, prefer it.
    try:
        from logger import get_logger  # type: ignore

        return get_logger()
    except Exception:
        logger = logging.getLogger("template_generator")
        if not logger.handlers:
            handler = logging.StreamHandler()
            formatter = logging.Formatter("%(levelname)s: %(message)s")
            handler.setFormatter(formatter)
            logger.addHandler(handler)
        logger.setLevel(logging.INFO)
        return logger


def _project_root() -> Path:
    return config.program_root()


def _col_letter(col_index_1_based: int) -> str:
    result = ""
    x = col_index_1_based
    while x > 0:
        x, rem = divmod(x - 1, 26)
        result = chr(ord("A") + rem) + result
    return result


def _set_headers_and_widths(ws, headers: list[str]) -> None:
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header).font = HEADER_FONT
        ws.cell(row=1, column=col_idx).alignment = HEADER_ALIGNMENT
        # Auto width based on header length.
        ws.column_dimensions[_col_letter(col_idx)].width = min(
            40, max(12, len(header) + 2)
        )

    ws.freeze_panes = FREEZE_PANES


def ensure_item_code_db() -> Path:
    """
    data/Item_Code_DB.xlsx 가 없을 때만 생성합니다. 이미 있으면 덮어쓰지 않습니다.
    """
    logger = _get_logger()
    d = config.data_dir()
    d.mkdir(parents=True, exist_ok=True)
    path = config.item_code_db_path()
    if path.exists():
        return path

    wb = Workbook()
    ws = wb.active
    ws.title = "Item_Code_DB"
    for col_idx, header in enumerate(ITEM_CODE_DB_HEADERS, start=1):
        ws.cell(row=1, column=col_idx, value=header).font = HEADER_FONT
        ws.cell(row=1, column=col_idx).alignment = HEADER_ALIGNMENT
        ws.column_dimensions[_col_letter(col_idx)].width = min(
            40, max(12, len(header) + 2)
        )
    ws.freeze_panes = "A2"
    for row_idx, row in enumerate(ITEM_CODE_DB_DEFAULT_ROWS, start=2):
        for col_idx, value in enumerate(row, start=1):
            ws.cell(row=row_idx, column=col_idx, value=value)
    wb.save(path)
    logger.info(f"Created Item_Code_DB: {path}")
    return path


def generate_class_define_template(
    output_path: Optional[Path | str] = None,
) -> Path:
    """
    Create `Class_Define_Template.xlsx` with required sheets:
    - Class_Define
    - Fluid_Service
    - Joint
    - Schedule
    - Reducing_Table
    - Pipe_Group
    - Fitting_Group
    - Flange
    - Valve

    동시에 data/Item_Code_DB.xlsx 가 없으면 생성합니다(기존 파일은 유지).
    """
    logger = _get_logger()

    template_path = (
        Path(output_path)
        if output_path is not None
        else _project_root() / DEFAULT_TEMPLATE_FILENAME
    )

    wb = Workbook()

    ws_define = wb.active
    ws_define.title = "Class_Define"
    _set_headers_and_widths(ws_define, CLASS_DEFINE_HEADERS)

    ws_fluid = wb.create_sheet(title="Fluid_Service")
    _set_headers_and_widths(ws_fluid, FLUID_SERVICE_HEADERS)

    ws_joint = wb.create_sheet(title="Joint")
    _set_headers_and_widths(ws_joint, JOINT_HEADERS)

    ws_schedule = wb.create_sheet(title="Schedule")
    _set_headers_and_widths(ws_schedule, SCHEDULE_HEADERS)

    ws_reducing_table = wb.create_sheet(title="Reducing_Table")
    _set_headers_and_widths(ws_reducing_table, REDUCING_TABLE_HEADERS)

    ws_pipe = wb.create_sheet(title="Pipe_Group")
    _set_headers_and_widths(ws_pipe, PIPE_HEADERS)

    ws_fitting = wb.create_sheet(title="Fitting_Group")
    _set_headers_and_widths(ws_fitting, FITTING_HEADERS)

    ws_flange = wb.create_sheet(title="Flange")
    _set_headers_and_widths(ws_flange, FLANGE_HEADERS)

    ws_valve = wb.create_sheet(title="Valve")
    _set_headers_and_widths(ws_valve, VALVE_HEADERS)

    try:
        wb.save(template_path)
    except PermissionError as e:
        raise PermissionError(
            f"엑셀 파일이 열려있어서 저장에 실패했습니다. "
            f"해당 파일을 닫고 다시 시도해 주세요. (path: {template_path})"
        ) from e

    ensure_item_code_db()

    logger.info(f"Generated template: {template_path}")
    return template_path


if __name__ == "__main__":
    generate_class_define_template()

